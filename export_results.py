import os
import sys
import json
import sqlite3
import argparse
from datetime import datetime, timezone
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "database.db")
RESULTS_DIR = os.path.join(BASE_DIR, "results")
DEFAULT_EXCEL_OUTPUT = os.path.join(RESULTS_DIR, "study_results.xlsx")


def get_local_tz_info():
    now_local = datetime.now().astimezone()
    tz_name = now_local.tzname() or "Local Time"
    utc_offset = now_local.strftime("%z")
    offset_formatted = f"UTC{utc_offset[:3]}:{utc_offset[3:]}" if utc_offset else "UTC+02:00"
    return now_local.tzinfo, tz_name, offset_formatted


def format_duration(seconds):
    if pd.isna(seconds) or seconds is None:
        return "N/A"
    try:
        total_secs = int(round(float(seconds)))
        mins = total_secs // 60
        secs = total_secs % 60
        return f"{mins:02d}:{secs:02d}"
    except (ValueError, TypeError):
        return "N/A"


def parse_and_convert_dates(series, local_tz):
    dt_utc = pd.to_datetime(series, utc=True, errors="coerce")
    dt_local = dt_utc.dt.tz_convert(local_tz)
    local_str = dt_local.dt.strftime("%Y-%m-%d %H:%M:%S").fillna("N/A")
    utc_str = dt_utc.dt.strftime("%Y-%m-%d %H:%M:%S").fillna("N/A")
    return local_str, utc_str


def build_kpi_summary(sessions_df, users_df):
    metrics = []
    total_u = len(users_df)
    comp = sessions_df[sessions_df["status"] == "commissioned"]
    aban = sessions_df[sessions_df["status"] == "abandoned"]

    metrics.append(("Total Registered Participants", total_u))
    metrics.append(("Commissioned Devices (Completed)", len(comp)))
    metrics.append(("Abandoned Sessions", len(aban)))

    comp_rate = (len(comp) / total_u * 100) if total_u else 0.0
    metrics.append(("Completion Rate (%)", f"{comp_rate:.1f}%"))

    avg_sec = comp["final_security_score"].mean() if not comp.empty else 0.0
    avg_use = comp["final_usability_score"].mean() if not comp.empty else 0.0
    metrics.append(("Mean Security Health Score (0-100)", f"{avg_sec:.1f}%"))
    metrics.append(("Mean Usability / Convenience Index (0-100)", f"{avg_use:.1f}%"))

    # Security Grade Distribution
    grade_a = len(comp[comp["security_grade"] == "A"])
    grade_b = len(comp[comp["security_grade"] == "B"])
    grade_c = len(comp[comp["security_grade"] == "C"])
    grade_d = len(comp[comp["security_grade"] == "D"])
    metrics.append(("Grade A (Hardened Zero-Trust)", f"{grade_a} ({grade_a/len(comp)*100:.1f}%)" if len(comp) else "0"))
    metrics.append(("Grade B (Balanced Consumer)", f"{grade_b} ({grade_b/len(comp)*100:.1f}%)" if len(comp) else "0"))
    metrics.append(("Grade C/D (High Vulnerability / Low Security)", f"{grade_c+grade_d} ({(grade_c+grade_d)/len(comp)*100:.1f}%)" if len(comp) else "0"))

    # Duration & Interactions
    avg_dur = comp["total_time_seconds"].mean() if not comp.empty else 0.0
    avg_toggles = comp["toggle_count"].mean() if not comp.empty else 0.0
    avg_reversals = comp["warning_reversals"].mean() if not comp.empty else 0.0

    metrics.append(("Average Configuration Time", format_duration(avg_dur)))
    metrics.append(("Average Settings Adjusted (Toggles Flipped)", f"{avg_toggles:.1f}"))
    metrics.append(("Average Warning Reversals (Nudge Rate)", f"{avg_reversals:.2f}"))

    return pd.DataFrame(metrics, columns=["Usability vs Security Metric", "Value"])


def load_database_data(db_path, local_tz, offset_label):
    if not os.path.exists(db_path):
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    conn = sqlite3.connect(db_path)

    users_df = pd.read_sql_query(
        "SELECT id AS user_id, email_hash AS participant_hash, created_at FROM users ORDER BY id ASC",
        conn
    )
    sessions_df = pd.read_sql_query(
        "SELECT * FROM commission_sessions ORDER BY id ASC",
        conn
    )
    logs_df = pd.read_sql_query(
        "SELECT * FROM interaction_logs ORDER BY id ASC",
        conn
    )
    conn.close()

    kpi_df = build_kpi_summary(sessions_df, users_df)

    if not users_df.empty:
        u_loc, u_utc = parse_and_convert_dates(users_df["created_at"], local_tz)
        users_df["registered_local"] = u_loc
        users_df["registered_utc"] = u_utc

    if not sessions_df.empty:
        s_start_loc, s_start_utc = parse_and_convert_dates(sessions_df["started_at"], local_tz)
        s_end_loc, s_end_utc = parse_and_convert_dates(sessions_df["completed_at"], local_tz)
        sessions_df["started_local"] = s_start_loc
        sessions_df["completed_local"] = s_end_loc
        sessions_df["duration_formatted"] = sessions_df["total_time_seconds"].apply(format_duration)

    return users_df, sessions_df, logs_df, kpi_df


def export_database_to_excel(output_path=None):
    os.makedirs(RESULTS_DIR, exist_ok=True)
    local_tz, tz_name, offset_label = get_local_tz_info()
    target_file = output_path or DEFAULT_EXCEL_OUTPUT

    users_df, sessions_df, logs_df, kpi_df = load_database_data(DB_PATH, local_tz, offset_label)

    # Style workbook with openpyxl
    wb = openpyxl.Workbook()

    # 1. Executive Summary Sheet
    ws_exec = wb.active
    ws_exec.title = "Executive_Summary"
    ws_exec.views.sheetView[0].showGridLines = True

    # Styling colors
    NAVY_DARK = "1E293B"
    ROYAL_BLUE = "2563EB"
    SLATE_LIGHT = "F1F5F9"
    WHITE = "FFFFFF"
    BORDER_COLOR = "CBD5E1"

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )

    # Title block
    ws_exec.merge_cells("B2:D2")
    ws_exec["B2"] = "IoT Usability vs. Security Trade-Off Research Analytics"
    ws_exec["B2"].font = Font(name="Segoe UI", size=14, bold=True, color=WHITE)
    ws_exec["B2"].fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    ws_exec["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_exec.row_dimensions[2].height = 36

    ws_exec.merge_cells("B3:D3")
    ws_exec["B3"] = "Commissioning Auditor Dataset | HaloGuard Framework"
    ws_exec["B3"].font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    ws_exec["B3"].fill = PatternFill(start_color=ROYAL_BLUE, end_color=ROYAL_BLUE, fill_type="solid")
    ws_exec["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws_exec.row_dimensions[3].height = 22

    # KPI Table Header
    ws_exec["B5"] = "Usability vs Security Metric"
    ws_exec["B5"].font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    ws_exec["B5"].fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    ws_exec["B5"].border = thin_border

    ws_exec.merge_cells("C5:D5")
    ws_exec["C5"] = "Study Benchmark Value"
    ws_exec["C5"].font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    ws_exec["C5"].fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    ws_exec["C5"].alignment = Alignment(horizontal="center", vertical="center")
    ws_exec["C5"].border = thin_border
    ws_exec["D5"].border = thin_border

    # KPI Data Rows
    row_idx = 6
    for _, row in kpi_df.iterrows():
        ws_exec.row_dimensions[row_idx].height = 22
        cell_m = ws_exec.cell(row=row_idx, column=2, value=row["Usability vs Security Metric"])
        cell_m.font = Font(name="Segoe UI", size=10, bold=("Mean" in str(row["Usability vs Security Metric"])))
        cell_m.border = thin_border

        ws_exec.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
        cell_v = ws_exec.cell(row=row_idx, column=3, value=row["Value"])
        cell_v.font = Font(name="Segoe UI", size=10)
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = thin_border
        ws_exec.cell(row=row_idx, column=4).border = thin_border
        row_idx += 1

    ws_exec.column_dimensions["B"].width = 45
    ws_exec.column_dimensions["C"].width = 25
    ws_exec.column_dimensions["D"].width = 25

    # 2. Participant Audits Sheet
    ws_audits = wb.create_sheet(title="Participant_Audits")
    ws_audits.views.sheetView[0].showGridLines = True
    audit_cols = [
        "Session ID", "User ID", "Status", "Security Score (%)", "Usability Index (%)",
        "Grade", "Risk Tier", "Duration (MM:SS)", "Toggles Adjusted", "Warning Reversals"
    ]
    ws_audits.append(audit_cols)
    ws_audits.row_dimensions[1].height = 26
    for cell in ws_audits[1]:
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if not sessions_df.empty:
        for _, s in sessions_df.iterrows():
            ws_audits.append([
                s.get("id"),
                f"P-{s.get('user_id', 1):03d}",
                s.get("status"),
                s.get("final_security_score"),
                s.get("final_usability_score"),
                s.get("security_grade"),
                s.get("risk_tier"),
                s.get("duration_formatted", "N/A"),
                s.get("toggle_count", 0),
                s.get("warning_reversals", 0),
            ])

    # 3. Setting Choices Breakdown Sheet
    ws_choices = wb.create_sheet(title="Setting_Choices")
    ws_choices.views.sheetView[0].showGridLines = True
    setting_cols = [
        "User ID", "Authentication", "Transport Crypto", "Network Topology",
        "Remote Exposure", "Firmware Patching", "Cloud Privacy", "Action Authorization"
    ]
    ws_choices.append(setting_cols)
    ws_choices.row_dimensions[1].height = 26
    for cell in ws_choices[1]:
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if not sessions_df.empty:
        for _, s in sessions_df.iterrows():
            ws_choices.append([
                f"P-{s.get('user_id', 1):03d}",
                s.get("auth_credential"),
                s.get("transport_crypto"),
                s.get("network_topology"),
                s.get("remote_exposure"),
                s.get("firmware_policy"),
                s.get("data_privacy"),
                s.get("action_auth"),
            ])

    # Auto-width for sheets
    for ws in [ws_audits, ws_choices]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(target_file)
    print(f"Report successfully saved to: {target_file}")
    return target_file


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export IoT Usability vs Security Research to Excel.")
    parser.add_argument("--output", type=str, default=None, help="Custom output path")
    args = parser.parse_args()
    export_database_to_excel(output_path=args.output)
